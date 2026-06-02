from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph


THESIS_DIR = Path(r"D:\计算机本科毕业论文")
DOCX_PATH = THESIS_DIR / "论文终极实验版.docx"
BACKUP_DOCX = THESIS_DIR / "论文终极实验版_修改前备份.docx"
ASSET_DIR = THESIS_DIR / "thesis_runtime_assets_v2"

MULTI_FORMAT_SRC_1 = Path(
    r"C:\Users\19892\Documents\Tencent Files\1989222763\nt_qq\nt_data\Pic\2026-05\Ori\048f6fe3548094914735988ed479a207.png"
)
MULTI_FORMAT_SRC_2 = Path(
    r"C:\Users\19892\Documents\Tencent Files\1989222763\nt_qq\nt_data\Pic\2026-05\Ori\f7a591e40ba6a09c8101259f1c246dd8.png"
)

INTERFACE_SRC_1 = Path(
    r"C:\Users\19892\Documents\Tencent Files\1989222763\nt_qq\nt_data\Pic\2026-05\Ori\609a3ab9d6265258c21975058f33235a.png"
)
INTERFACE_SRC_2 = Path(
    r"C:\Users\19892\Documents\Tencent Files\1989222763\nt_qq\nt_data\Pic\2026-05\Ori\be099b16186188dbecc8116f74a18199.png"
)
INTERFACE_SRC_3 = Path(
    r"C:\Users\19892\Documents\Tencent Files\1989222763\nt_qq\nt_data\Pic\2026-05\Ori\075066939682fd3a5bbbc2d57d7ab3de.png"
)

CASE_SRC_1 = Path(
    r"C:\Users\19892\Documents\Tencent Files\1989222763\nt_qq\nt_data\Pic\2026-05\Ori\c4335118575c27b936d9c1b982feebb7.png"
)
CASE_SRC_2 = Path(
    r"C:\Users\19892\Documents\Tencent Files\1989222763\nt_qq\nt_data\Pic\2026-05\Ori\342f718456af5d7b0529eb7b6c00b8b1.png"
)
CASE_SRC_3 = Path(
    r"C:\Users\19892\Documents\Tencent Files\1989222763\nt_qq\nt_data\Pic\2026-05\Ori\7199c9d389e7110c15db6b023fe6a57a.png"
)
CASE_SRC_4 = Path(
    r"C:\Users\19892\Documents\Tencent Files\1989222763\nt_qq\nt_data\Pic\2026-05\Ori\c9fa629388f0d4219d76b3b72fb759b1.png"
)
CASE_SRC_5 = Path(
    r"C:\Users\19892\Documents\Tencent Files\1989222763\nt_qq\nt_data\Pic\2026-05\Ori\b396e3c79600a6fdba9d2965a845c4c5.png"
)

FIG_4_4 = ASSET_DIR / "figure_4_4_multiformat.png"
FIG_4_6 = ASSET_DIR / "figure_4_6_runtime_upload.png"
FIG_4_7 = ASSET_DIR / "figure_4_7_history_correction.png"
FIG_5_6 = ASSET_DIR / "figure_5_6_runtime_interfaces.png"
FIG_5_7 = ASSET_DIR / "figure_5_7_real_cases.png"


def backup_docx() -> None:
    if not BACKUP_DOCX.exists():
        shutil.copy2(DOCX_PATH, BACKUP_DOCX)


def load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def fit_contain(image: Image.Image, max_width: int, max_height: int) -> Image.Image:
    scale = min(max_width / image.width, max_height / image.height)
    new_size = (max(1, int(image.width * scale)), max(1, int(image.height * scale)))
    return image.resize(new_size, Image.LANCZOS)


def paste_panel(
    canvas: Image.Image,
    source: Path,
    box: tuple[int, int, int, int],
    label: str,
    font: ImageFont.ImageFont,
) -> None:
    img = Image.open(source).convert("RGB")
    x0, y0, x1, y1 = box
    pad = 20
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle(box, radius=18, outline="#cfd8e5", width=2, fill="white")
    label_x = x0 + 16
    label_y = y0 + 10
    draw.text((label_x, label_y), label, fill="#1d3557", font=font)
    content = fit_contain(img, x1 - x0 - pad * 2, y1 - y0 - pad * 2 - 40)
    paste_x = x0 + (x1 - x0 - content.width) // 2
    paste_y = y0 + 48 + (y1 - y0 - 48 - content.height) // 2
    canvas.paste(content, (paste_x, paste_y))


def build_figures() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    label_font = load_font(36)

    # Figure 4.4
    canvas = Image.new("RGB", (2300, 2800), "white")
    paste_panel(canvas, MULTI_FORMAT_SRC_1, (60, 60, 2240, 1380), "(a)", label_font)
    paste_panel(canvas, MULTI_FORMAT_SRC_2, (60, 1460, 2240, 2740), "(b)", label_font)
    canvas.save(FIG_4_4, quality=95)

    # Figure 4.6
    img = Image.open(INTERFACE_SRC_1).convert("RGB")
    img.save(FIG_4_6, quality=95)

    # Figure 4.7
    canvas = Image.new("RGB", (2200, 2500), "white")
    paste_panel(canvas, INTERFACE_SRC_2, (60, 60, 2140, 1450), "(a)", label_font)
    paste_panel(canvas, INTERFACE_SRC_3, (260, 1540, 1940, 2440), "(b)", label_font)
    canvas.save(FIG_4_7, quality=95)

    # Figure 5.6
    canvas = Image.new("RGB", (2400, 2700), "white")
    paste_panel(canvas, INTERFACE_SRC_1, (60, 60, 1160, 1260), "(a)", label_font)
    paste_panel(canvas, INTERFACE_SRC_2, (1240, 60, 2340, 1260), "(b)", label_font)
    paste_panel(canvas, INTERFACE_SRC_3, (420, 1360, 1980, 2600), "(c)", label_font)
    canvas.save(FIG_5_6, quality=95)

    # Figure 5.7
    canvas = Image.new("RGB", (2400, 3300), "white")
    paste_panel(canvas, CASE_SRC_1, (60, 60, 1160, 1060), "(a)", label_font)
    paste_panel(canvas, CASE_SRC_2, (1240, 60, 2340, 1060), "(b)", label_font)
    paste_panel(canvas, CASE_SRC_3, (60, 1140, 1160, 2140), "(c)", label_font)
    paste_panel(canvas, CASE_SRC_4, (1240, 1140, 2340, 2140), "(d)", label_font)
    paste_panel(canvas, CASE_SRC_5, (360, 2220, 2040, 3220), "(e)", label_font)
    canvas.save(FIG_5_7, quality=95)


def find_paragraph_exact(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def paragraph_index(doc: Document, paragraph: Paragraph) -> int:
    target = paragraph._p
    for index, item in enumerate(doc.paragraphs):
        if item._p is target:
            return index
    raise ValueError("Paragraph index not found.")


def next_paragraph(doc: Document, paragraph: Paragraph, offset: int = 1) -> Paragraph:
    paragraphs = doc.paragraphs
    index = paragraph_index(doc, paragraph)
    return paragraphs[index + offset]


def previous_paragraph(doc: Document, paragraph: Paragraph) -> Paragraph:
    paragraphs = doc.paragraphs
    index = paragraph_index(doc, paragraph)
    return paragraphs[index - 1]


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag.endswith("}pPr"):
            continue
        p.remove(child)


def set_paragraph_text(paragraph: Paragraph, text: str) -> Paragraph:
    alignment = paragraph.alignment
    style = paragraph.style
    clear_paragraph(paragraph)
    paragraph.style = style
    paragraph.alignment = alignment
    paragraph.add_run(text)
    return paragraph


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = paragraph.style
    if text:
        new_para.add_run(text)
    return new_para


def add_center_picture(paragraph: Paragraph, image_path: Path, width_cm: float) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))


def set_center_caption(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(text)


def update_table(table, rows: list[list[str]]) -> None:
    max_cols = max(len(row) for row in rows)
    while len(table.columns) < max_cols:
        table.add_column(Cm(3.0))
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            table.cell(r_idx, c_idx).text = value


def normalize_heading_paragraphs(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name
        if style_name in {"Heading 1", "Heading 2", "Heading 3", "标题 1", "标题 2", "标题 3"}:
            text = paragraph.text
            if not text.strip():
                continue
            alignment = paragraph.alignment
            style = paragraph.style
            clear_paragraph(paragraph)
            paragraph.style = style
            paragraph.alignment = alignment
            paragraph.add_run(text)


def update_body_texts(doc: Document) -> None:
    replacements = {
        "结合项目实现，系统功能可归纳为用户管理、公式输入、识别与解析、结果展示导出、历史管理以及勘误反馈六个方面。": "结合项目实现，系统功能可归纳为用户管理、公式输入、识别与解析、结果展示导出、历史管理以及勘误反馈六个方面。其中，在线手写板功能直接影响系统的课堂交互体验，而指笔混合输入相关研究说明，自然书写输入的效率和反馈及时性会显著影响用户对系统的接受程度[57]。相比仅提供单次识别结果的原型系统，当前版本进一步把纠错纳入核心功能范围，使系统在识别出错时仍能维持可用性。此外，系统还需要支持识别结果的多格式导出和勘误知识积累，即同一条识别结果应能稳定转换为 LaTeX、Word 标准公式和 PNG 渲染图，同时把用户提交的纠错持续沉淀为可复用的候选知识。",
        "结合项目实现过程可以看到，本文的研究思路并不是把“模型训练”和“系统开发”割裂开来，而是围绕课堂公式录入这一具体场景，把输入采集、图像预处理、序列识别、多格式输出、历史留存与人工纠错组织为同一条工作链路。与只关注离线精度的纯算法实验相比，这种架构更强调前后端协同和结果可复用性，也更贴合本科毕业设计对可展示性与完整性的要求。如图 1.1 所示。": "结合项目实现过程可以看到，本文的研究思路并不是把“模型训练”和“系统开发”割裂开来，而是围绕课堂公式录入这一具体场景，把输入采集、图像预处理、序列识别、多格式输出、历史留存与人工纠错组织为同一条工作链路。与只关注离线精度的纯算法实验相比，这种架构更强调前后端协同和结果可复用性，也更贴合本科毕业设计对可展示性与完整性的要求，其总体技术路线如图 1.1 所示。",
        "全文共分为六章。第一章介绍课题背景、研究现状以及本文的总体技术路线；第二章归纳手写数学公式识别涉及的相关理论与关键技术；第三章从需求分析、系统角色与模块划分等角度给出总体设计；第四章结合项目代码，对前端输入、预处理、识别模型、结果转换、历史记录与勘误闭环等关键实现进行详细说明；第五章基于训练日志、实际运行界面和项目数据文件，对模型性能、系统功能和真实样例进行综合测试与分析；第六章总结全文工作，并讨论后续仍可继续完善的方向。如图 1.2 所示。": "全文共分为六章。第一章介绍课题背景、研究现状以及本文的总体技术路线；第二章归纳手写数学公式识别涉及的相关理论与关键技术；第三章从需求分析、系统角色与模块划分等角度给出总体设计；第四章结合项目代码，对前端输入、预处理、识别模型、结果转换、历史记录与勘误闭环等关键实现进行详细说明；第五章基于训练日志、实际运行界面和项目数据文件，对模型性能、系统功能和真实样例进行综合测试与分析；第六章总结全文工作，并讨论后续仍可继续完善的方向。论文结构安排总体框架如图 1.2 所示。",
        "系统整体采用基于 Flask 的 Web 应用架构。前端负责用户交互、文件上传、画布书写和结果展示；后端负责路由分发、身份校验、图像预处理、模型调用、格式转换和历史记录保存。结合项目代码可知，主流程已经形成相对清晰的“输入层—推理层—结果层”分层结构。如图 3.1 所示。": "系统整体采用基于 Flask 的 Web 应用架构。前端负责用户交互、文件上传、画布书写和结果展示；后端负责路由分发、身份校验、图像预处理、模型调用、格式转换和历史记录保存。结合项目代码可知，主流程已经形成相对清晰的“输入层—推理层—结果层”分层结构，其总体架构如图 3.1 所示。",
        "当用户提交识别请求后，系统首先判断输入来源并完成统一解码。随后执行预处理，把图像变换为模型可接受的二值化画布，再调用识别模型输出 token 序列。最后，系统将序列结果转换为可展示、可复制和可计算的表达形式，并写入历史记录。若用户发现识别结果存在问题，则可进一步进入勘误流程：前端弹出候选修正窗口，后端基于已有勘误表和内置公式种子推荐若干候选项；用户选择候选结果或手动输入正确公式后，系统会同步更新历史记录中的当前识别结果、保留原识别文本，并将本次修正写入勘误表，供后续相似错误复用。如图 3.2 所示。如图 3.3 所示。如图 3.4 所示。": "当用户提交识别请求后，系统首先判断输入来源并完成统一解码。随后执行预处理，把图像变换为模型可接受的二值化画布，再调用识别模型输出 token 序列。最后，系统将序列结果转换为可展示、可复制和可计算的表达形式，并写入历史记录。若用户发现识别结果存在问题，则可进一步进入勘误流程：前端弹出候选修正窗口，后端基于已有勘误表和内置公式种子推荐若干候选项；用户选择候选结果或手动输入正确公式后，系统会同步更新历史记录中的当前识别结果、保留原识别文本，并将本次修正写入勘误表，供后续相似错误复用。整体业务流程如图 3.2 所示，前后端交互时序如图 3.3 所示，勘误闭环过程如图 3.4 所示。",
        "在噪声处理中，系统使用连通域面积过滤去除过小噪声块，再通过边界扩展和前景居中操作保留主要书写区域。对于笔画像素占比较低的样本，代码中还加入了按比例控制的轻度膨胀步骤，用以增强细弱笔画，这与非规范手写样式研究中强调的“保持结构可辨识度”思路相一致[53]。如表 4.1 所示。如图 4.1 所示。": "在噪声处理中，系统使用连通域面积过滤去除过小噪声块，再通过边界扩展和前景居中操作保留主要书写区域。对于笔画像素占比较低的样本，代码中还加入了按比例控制的轻度膨胀步骤，用以增强细弱笔画，这与非规范手写样式研究中强调的“保持结构可辨识度”思路相一致[53]。预处理步骤及作用说明见表 4.1，图像预处理流程如图 4.1 所示。",
        "本文没有采用传统字符切分方案，而是将整幅公式图像直接输入模型，由网络输出 token 序列。这样做的好处在于训练目标统一，避免了切分误差向后传播，也更适合处理分数、根式和上下标等二维结构明显的公式[13]。如图 4.2 所示。": "本文没有采用传统字符切分方案，而是将整幅公式图像直接输入模型，由网络输出 token 序列。这样做的好处在于训练目标统一，避免了切分误差向后传播，也更适合处理分数、根式和上下标等二维结构明显的公式[13]。端到端识别模型结构如图 4.2 所示。",
        "从 train/dataset_e2e_v3.py 和 train/train_e2e_v3.py 可以看到，V3 训练流程加入了更强的数据增强策略，包括弹性形变、仿射变换、噪声注入、随机擦除、CutOut 和笔画粗细变化等。这类策略的目标不是简单“增加样本数量”，而是尽量模拟真实手写中的姿态变化和拍照扰动。树式数据增强和语法数据生成等研究表明，面向公式结构的样本扩展能够有效缓解样本不足和结构偏差问题[25][26]。同时，项目训练脚本采用 AdamW 优化器与分阶段学习率调度，其自适应梯度更新思想可追溯到 Adam 方法[40]。国内关于在线手写公式合成的研究也说明，合理的数据生成和增强有助于缓解公式样本不足问题[49]。如图 4.3 所示。": "从 train/dataset_e2e_v3.py 和 train/train_e2e_v3.py 可以看到，V3 训练流程加入了更强的数据增强策略，包括弹性形变、仿射变换、噪声注入、随机擦除、CutOut 和笔画粗细变化等。这类策略的目标不是简单“增加样本数量”，而是尽量模拟真实手写中的姿态变化和拍照扰动。树式数据增强和语法数据生成等研究表明，面向公式结构的样本扩展能够有效缓解样本不足和结构偏差问题[25][26]。同时，项目训练脚本采用 AdamW 优化器与分阶段学习率调度，其自适应梯度更新思想可追溯到 Adam 方法[40]。国内关于在线手写公式合成的研究也说明，合理的数据生成和增强有助于缓解公式样本不足问题[49]。关键改进技术原理如图 4.3 所示。",
        "模型性能仍以 sequence accuracy 和 token accuracy 为主。前者要求整条公式完全正确，更接近真实使用时“能否直接拿去用”的判断标准；后者反映局部符号预测质量，适合观察模型在细粒度层面的提升情况。除此之外，本文还增加了两类工程化验证口径：一类是系统功能是否能够完整跑通，包括登录、上传、手写、导出、历史回看与勘误回写；另一类是项目数据文件中已经积累的运行记录，它能够从侧面反映系统在不同输入方式和纠错场景下的实际使用特征[6]。如图 5.1 所示。": "模型性能仍以 sequence accuracy 和 token accuracy 为主。前者要求整条公式完全正确，更接近真实使用时“能否直接拿去用”的判断标准；后者反映局部符号预测质量，适合观察模型在细粒度层面的提升情况。除此之外，本文还增加了两类工程化验证口径：一类是系统功能是否能够完整跑通，包括登录、上传、手写、导出、历史回看与勘误回写；另一类是项目数据文件中已经积累的运行记录，它能够从侧面反映系统在不同输入方式和纠错场景下的实际使用特征[6]。数据集划分与训练来源示意如图 5.1 所示。",
        "从项目保留的训练日志与检查点指标来看，V2 和 V3 模型的主体框架保持一致，差异主要集中在编码增强、对齐建模、数据增强和解码策略等几个关键环节。正因为改动具有明确边界，因此两版模型的对比更能说明“具体工程改进是否带来了可观收益”，而不是被其他无关变量干扰。如表 5.1 所示。": "从项目保留的训练日志与检查点指标来看，V2 和 V3 模型的主体框架保持一致，差异主要集中在编码增强、对齐建模、数据增强和解码策略等几个关键环节。正因为改动具有明确边界，因此两版模型的对比更能说明“具体工程改进是否带来了可观收益”，而不是被其他无关变量干扰，关键配置对比见表 5.1。",
        "结合训练日志整理得到的最佳指标显示，V2 模型的 sequence accuracy 为 70.56%，token accuracy 为 86.58%；V3 模型分别提升到 83.24% 和 92.69%。如果只看数值，二者差距已经比较明显；若结合项目场景理解，这种提升更重要的意义在于它减少了“公式大体看起来像对了，但关键结构仍然错位”的情况。如表 5.2 所示。如图 5.2 所示。如图 5.3 所示。": "结合训练日志整理得到的最佳指标显示，V2 模型的 sequence accuracy 为 70.56%，token accuracy 为 86.58%；V3 模型分别提升到 83.24% 和 92.69%。如果只看数值，二者差距已经比较明显；若结合项目场景理解，这种提升更重要的意义在于它减少了“公式大体看起来像对了，但关键结构仍然错位”的情况。相关性能对比见表 5.2，模型性能柱状对比和验证集序列准确率变化曲线分别如图 5.2 和图 5.3 所示。",
        "结合本次实际运行测试，可以确认登录认证、图片上传、在线手写、LaTeX/Word/PNG 导出、历史记录查看和勘误弹窗等核心功能均能正常工作。尤其是在上传二次公式样例后，系统可以同时返回公式渲染图、LaTeX 文本、Word Linear 结果以及多个勘误推荐项，这一过程与第四章对实现链路的描述能够互相印证。如表 5.3 所示。如图 5.4 所示。如图 5.5 所示。": "结合本次实际运行测试，可以确认登录认证、图片上传、在线手写、LaTeX/Word/PNG 导出、历史记录查看和勘误弹窗等核心功能均能正常工作。尤其是在上传二次公式样例后，系统可以同时返回公式渲染图、LaTeX 文本、Word Linear 结果以及多个勘误推荐项，这一过程与第四章对实现链路的描述能够互相印证。系统主要功能测试结果见表 5.3，识别记录来源分布和勘误知识表来源构成分别如图 5.4 和图 5.5 所示。",
    }

    for paragraph in doc.paragraphs:
        for prefix, new_text in replacements.items():
            if paragraph.text.startswith(prefix):
                set_paragraph_text(paragraph, new_text)
                break


def update_section_441(doc: Document) -> None:
    heading = find_paragraph_exact(doc, "4.4.1 多格式转换与公式渲染")
    body1 = next_paragraph(doc, heading)
    body2 = next_paragraph(doc, heading, 2)
    set_paragraph_text(
        body1,
        "模型输出的原始结果是离散 token 序列。系统首先按规则将其拼接为近似 LaTeX 的文本表达，再对幂运算、括号、分式符号和特殊字符进行规范化处理，以提高可读性和后续解析成功率。对于需要进入勘误表或历史记录的结果，后端还会调用 `_normalize_formula_for_storage()` 做统一存储处理，尽量把同义写法归并到稳定的内部表达形式。",
    )
    set_paragraph_text(
        body2,
        "在展示层，系统并不是只给出一条单一字符串，而是围绕不同使用场景同步生成三类输出结果：一类是便于再次编辑和论文排版的 LaTeX 文本；一类是便于直接粘贴到 Office 文档中的 Word 标准公式结果；另一类是便于页面预览、复制和快速分享的 PNG 渲染图。三种结果共用同一识别来源，但在编辑性、展示性和文档兼容性上各有侧重。",
    )

    caption = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "图 4.4 多格式输出示例":
            caption = paragraph
            break
    if caption is None:
        p_desc = insert_paragraph_after(body2)
        set_paragraph_text(
            p_desc,
            "以在线手写输入的 y=(4ac-b^2)/(4a) 为例，系统在识别完成后会同时返回 LaTeX 结果、Word 标准公式粘贴效果和 PNG 渲染图。对于用户而言，LaTeX 适合继续修改和学术排版，Word 结果适合直接嵌入课程文档或论文，PNG 结果则适合网页展示和快速转发，多格式输出示例如图 4.4 所示。",
        )
        fig_para = insert_paragraph_after(p_desc)
        add_center_picture(fig_para, FIG_4_4, 15.0)
        caption = insert_paragraph_after(fig_para)
    set_center_caption(caption, "图 4.4 多格式输出示例")


def update_section_452(doc: Document) -> None:
    heading = find_paragraph_exact(doc, "4.5.2 勘误知识表与人工回写")
    body1 = next_paragraph(doc, heading)
    body2 = next_paragraph(doc, heading, 2)
    body3 = next_paragraph(doc, heading, 3)

    set_paragraph_text(
        body1,
        "从数据库实现细节看，`record_correction()` 会把新的错误公式与正确公式配对写入 `correction_table`；如果同类修正已经存在，则只累计使用次数并刷新更新时间。与此同时，`update_history_correction()` 会把原识别结果转存到 `original_recognized_text`，再用修正后的公式覆盖 `recognized_text`。这样一来，系统既保留了“模型第一次给出的答案”，也保留了“用户最终确认的答案”，为后续分析常见错误来源提供了可追溯依据。",
    )
    set_paragraph_text(
        body2,
        "除历史记录外，项目还内置了 11 条常见教学公式勘误种子，覆盖二次方程、解析几何、圆与三角、数列以及三角恒等式等高频类型。内置种子的作用不是替代人工判断，而是在系统首次部署时就提供一批可直接调用的候选项，使用户在遇到典型错误时不必每次都从零开始手动改写。",
    )
    set_paragraph_text(
        body3,
        "当用户提交修正后，系统一方面会把当前错误—正确公式对持续沉淀到 `correction_table` 中，另一方面会同步回写对应历史记录并保留原识别文本。由此，整套机制形成了“内置种子提供初始候选—用户人工确认或手动输入—历史记录回写—勘误知识持续扩充”的闭环路径。勘误知识表的种子类型与闭环作用见表 4.2，JSON 存储结构示意如图 4.5 所示，系统输入与结果展示界面如图 4.6 所示，勘误推荐与历史记录界面如图 4.7 所示。",
    )

    table_caption = find_paragraph_exact(doc, "表 4.2 勘误表相关数据项及作用")
    set_center_caption(table_caption, "表 4.2 勘误知识表种子类型与闭环作用")
    update_table(
        doc.tables[2],
        [
            ["类别", "代表公式", "初始化方式", "闭环作用"],
            ["二次函数与方程类", "求根公式、判别式、顶点坐标", "系统内置种子", "覆盖课堂高频公式，降低首次纠错门槛"],
            ["解析几何类", "两点距离、斜率公式", "系统内置种子", "处理上下标、分式与坐标差结构错误"],
            ["圆与三角类", "圆面积、圆周长、三角恒等式", "系统内置种子", "支撑 π、θ 与三角函数相关公式纠错"],
            ["数列类", "等差数列求和公式", "系统内置种子", "处理下标、乘除关系和分式结构错误"],
            ["人工修正类", "用户提交的错误—正确公式对", "运行过程累积", "通过回写历史与使用次数统计持续扩充知识表"],
        ],
    )

    set_center_caption(find_paragraph_exact(doc, "图 4.4 数据库（JSON 存储）ER 示意图"), "图 4.5 数据库（JSON 存储）ER 示意图")
    set_center_caption(find_paragraph_exact(doc, "图 4.5 系统输入与结果展示界面"), "图 4.6 系统输入与结果展示界面")
    set_center_caption(find_paragraph_exact(doc, "图 4.6 勘误表推荐与历史记录界面"), "图 4.7 勘误表推荐与历史记录界面")

    fig_46_caption = find_paragraph_exact(doc, "图 4.6 系统输入与结果展示界面")
    fig_47_caption = find_paragraph_exact(doc, "图 4.7 勘误表推荐与历史记录界面")
    add_center_picture(previous_paragraph(doc, fig_46_caption), FIG_4_6, 15.0)
    add_center_picture(previous_paragraph(doc, fig_47_caption), FIG_4_7, 15.0)


def update_section_54(doc: Document) -> None:
    heading = find_paragraph_exact(doc, "5.4 真实样例与界面运行分析")
    intro = next_paragraph(doc, heading)
    discussion = find_paragraph_exact(
        doc,
        "从复测结果看，系统对一次函数、简单根式、标准二次公式以及常见三角面积公式已经表现出较好的可用性；真正容易失稳的，仍然是层次更深、分式嵌套更复杂、或者局部连笔比较明显的样例。换句话说，当前模型在“常见教学公式”层面已经具备展示价值，但在更复杂结构下仍需要依赖后续纠错或进一步训练来保证结果可靠。",
    )
    set_paragraph_text(
        intro,
        "为更接近答辩演示和真实使用场景，本文重新选取了 3 组现场测试样例，并补充了系统真实运行界面截图进行说明。界面部分包括图片上传与结果展示页、历史记录页和勘误弹窗；样例部分包括球的体积公式、等差数列求和公式和三角恒等式。其中，等差数列样例在首次识别时出现结构性错误，因此保留了完整的人工纠错过程，用于展示勘误闭环在真实使用中的工作方式。系统真实运行界面如图 5.6 所示，三组样例及纠错过程如图 5.7 所示，测试结果汇总见表 5.4。",
    )
    set_paragraph_text(
        discussion,
        "从三次复测结果看，球的体积公式和三角恒等式都能被系统较完整地识别，说明当前模型对结构相对清晰、符号关系稳定的常见教学公式已经具备较好的可用性。相比之下，等差数列求和公式在初次识别时把分母 2 误判为 n，导致公式结构虽大体接近目标形式，但关键分式语义发生偏移；不过用户可以立即通过勘误弹窗手动输入正确公式，系统随后同步更新当前结果并保留原识别文本。这表明当前系统虽然还不能保证所有公式一次识别正确，但已经具备“结果可复核、错误可回写、经验可积累”的工程闭环，这种设计比单纯追求一次性输出更符合课堂和作业录入场景的实际需求。",
    )

    fig_56_caption = find_paragraph_exact(doc, "图 5.6 真实样例现场测试结果图")
    fig_57_caption = find_paragraph_exact(doc, "图 5.7 真实样例人工评估统计图")
    table_caption = find_paragraph_exact(doc, "表 5.4 真实样例现场测试结果汇总")

    add_center_picture(previous_paragraph(doc, fig_56_caption), FIG_5_6, 15.0)
    set_center_caption(fig_56_caption, "图 5.6 系统真实运行界面")

    update_table(
        doc.tables[6],
        [
            ["样例", "测试公式", "初次识别情况", "纠错情况", "最终评价"],
            ["球体积公式", "S=4/3 πr^3", "分式与幂次结构识别正确，可正常生成 LaTeX、Word 和 PNG 结果", "无需纠错", "正确"],
            ["等差数列求和公式", "S_n=na_1+n(n-1)d/2", "初次把分母 2 误识别为 n，整体结构接近但关键分式语义错误", "通过勘误弹窗手动输入正确公式并回写历史", "纠错后正确"],
            ["三角恒等式", "sin^2θ+cos^2θ=1", "三角函数与幂次关系识别较完整，恒等关系正确保留", "无需纠错", "正确"],
        ],
    )
    set_center_caption(table_caption, "表 5.4 三组真实公式样例测试结果汇总")

    add_center_picture(previous_paragraph(doc, fig_57_caption), FIG_5_7, 14.6)
    set_center_caption(fig_57_caption, "图 5.7 三组真实公式样例测试与纠错过程")


def main() -> None:
    backup_docx()
    build_figures()
    doc = Document(str(DOCX_PATH))
    normalize_heading_paragraphs(doc)
    update_body_texts(doc)
    update_section_441(doc)
    update_section_452(doc)
    update_section_54(doc)
    doc.save(str(DOCX_PATH))


if __name__ == "__main__":
    main()

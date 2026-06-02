from __future__ import annotations

import json
import math
import re
import shutil
from collections import Counter
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


PROJECT_DIR = Path(r"D:\Handwritten-Mathematical-Equation-Recognition-Using-CNN-main")
THESIS_DIR = Path(r"D:\计算机本科毕业论文")
BASE_DOC = THESIS_DIR / "论文终极版.docx"
SOURCE_DOC = THESIS_DIR / "论文3.4_修订稿_参考文献规范版.docx"
OUTPUT_DOC = THESIS_DIR / "论文4.0_终稿优化版.docx"
NOTE_FILE = THESIS_DIR / "论文4.0_修改说明.md"
FIG_DIR = THESIS_DIR / "generated_figures_v4"
UI_DIR = FIG_DIR / "ui"
V2_FIG_DIR = THESIS_DIR / "generated_figures_v2"

CN_FONT = Path(r"C:\Windows\Fonts\msyh.ttc")
CN_BOLD_FONT = Path(r"C:\Windows\Fonts\msyhbd.ttc")
SERIF_FONT = Path(r"C:\Windows\Fonts\simsun.ttc")


def ensure_dirs() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    UI_DIR.mkdir(parents=True, exist_ok=True)


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = CN_BOLD_FONT if bold and CN_BOLD_FONT.exists() else CN_FONT
    if not path.exists():
        path = SERIF_FONT
    return ImageFont.truetype(str(path), size=size)


def plot_font(size: int = 12) -> font_manager.FontProperties:
    font_path = CN_FONT if CN_FONT.exists() else SERIF_FONT
    return font_manager.FontProperties(fname=str(font_path), size=size)


def load_runtime_stats() -> dict:
    store = json.loads((PROJECT_DIR / "data" / "classroom_app.json").read_text(encoding="utf-8"))
    history = store["recognition_history"]
    corrections = store.get("correction_table", [])
    return {
        "history_total": len(history),
        "history_source": Counter(item.get("source_type", "") for item in history),
        "correction_total": len(corrections),
        "correction_source": Counter(item.get("source_type", "") for item in corrections),
    }


def load_training_metrics() -> dict:
    metrics = {}
    for filename, label in [("train_e2e_log.csv", "V2"), ("train_e2e_v3_log.csv", "V3")]:
        rows = []
        for line in (PROJECT_DIR / "logs" / filename).read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("epoch,train_loss"):
                continue
            parts = line.split(",")
            if not parts[0].isdigit():
                continue
            rows.append(
                {
                    "epoch": int(parts[0]),
                    "val_seq_acc": float(parts[5]),
                    "val_tok_acc_decode": float(parts[6]),
                }
            )
        rows.sort(key=lambda item: item["epoch"])
        metrics[label] = rows
    return metrics


def load_real_cases() -> list[dict]:
    path = V2_FIG_DIR / "real_case_results.json"
    if not path.exists():
        return []
    return json.loads(path.read_text(encoding="utf-8"))


def rounded_rect(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], fill: str, outline: str, radius: int = 26) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=2)


def fit_image(img: Image.Image, size: tuple[int, int]) -> Image.Image:
    canvas = Image.new("RGB", size, "#ffffff")
    copy = img.copy()
    copy.thumbnail((size[0] - 28, size[1] - 28))
    canvas.paste(copy, ((size[0] - copy.width) // 2, (size[1] - copy.height) // 2))
    return canvas


def ui_card(canvas: Image.Image, box: tuple[int, int, int, int], title: str, badge: str, image_path: Path, fill: str) -> None:
    draw = ImageDraw.Draw(canvas)
    rounded_rect(draw, box, fill, "#d7e1ec")
    badge_box = (box[0] + 26, box[1] + 22, box[0] + 180, box[1] + 62)
    rounded_rect(draw, badge_box, "#ffffff", "#d7e1ec", 18)
    draw.text((badge_box[0] + 18, badge_box[1] + 8), badge, fill="#264f78", font=pil_font(22, bold=True))
    draw.text((box[0] + 26, box[1] + 82), title, fill="#173a61", font=pil_font(30, bold=True))
    image_box = (box[0] + 24, box[1] + 126, box[2] - 24, box[3] - 24)
    rounded_rect(draw, image_box, "#ffffff", "#dce5ef", 22)
    image = Image.open(image_path).convert("RGB")
    framed = fit_image(image, (image_box[2] - image_box[0], image_box[3] - image_box[1]))
    canvas.paste(framed, (image_box[0], image_box[1]))


def make_architecture_figure() -> Path:
    img = Image.new("RGB", (1800, 980), "#f6f8fc")
    draw = ImageDraw.Draw(img)
    draw.text((575, 36), "系统总体架构与技术路线图", fill="#17395f", font=pil_font(40, bold=True))
    draw.text((475, 90), "从课堂输入、统一预处理、序列识别、多格式输出到勘误积累，形成面向教学应用的完整闭环", fill="#6d829b", font=pil_font(21))

    boxes = [
        ((76, 228, 390, 468), "#eaf4ff", "#7ba8d6", "应用场景层", ["课堂拍照上传", "在线手写输入", "教师与学生复用"]),
        ((424, 178, 770, 418), "#eefaf3", "#71b393", "Web 交互层", ["登录注册与角色校验", "拖拽上传与 Canvas 采集", "结果展示与复制导出"]),
        ((834, 178, 1180, 418), "#fff6ea", "#dfa26c", "后端处理层", ["图像解码与通道统一", "灰度化、阈值分割、裁剪", "接口组织与结果回传"]),
        ((1244, 178, 1608, 418), "#f5f0ff", "#ab8fda", "识别与转换层", ["CNN + BiGRU + Decoder", "Beam Search 与注意力对齐", "LaTeX / Word / MathML"]),
        ((724, 560, 1232, 840), "#eef2ff", "#7d96d2", "数据积累与反馈层", ["CROHME 训练集与模型日志", "识别历史记录与导出文件", "勘误知识表与人工修正回写"]),
    ]
    for box, fill, outline, title, lines in boxes:
        rounded_rect(draw, box, fill, outline)
        draw.text((box[0] + 24, box[1] + 24), title, fill="#1c3957", font=pil_font(28, bold=True))
        for idx, line in enumerate(lines):
            draw.text((box[0] + 24, box[1] + 78 + idx * 42), line, fill="#4f667f", font=pil_font(19))

    arrows = [
        ((390, 350), (424, 302)),
        ((770, 302), (834, 302)),
        ((1180, 302), (1244, 302)),
        ((600, 418), (910, 560)),
        ((1006, 418), (1006, 560)),
        ((1416, 418), (1120, 560)),
    ]
    for start, end in arrows:
        draw.line([start, end], fill="#355d8b", width=5)
        ang = math.atan2(end[1] - start[1], end[0] - start[0])
        size = 16
        p1 = (end[0] - size * math.cos(ang - math.pi / 6), end[1] - size * math.sin(ang - math.pi / 6))
        p2 = (end[0] - size * math.cos(ang + math.pi / 6), end[1] - size * math.sin(ang + math.pi / 6))
        draw.polygon([end, p1, p2], fill="#355d8b")

    out = FIG_DIR / "figure_1_architecture.png"
    img.save(out, quality=95)
    return out


def make_ui_collages() -> dict:
    result = {}

    canvas = Image.new("RGB", (1800, 980), "#f7f9fc")
    draw = ImageDraw.Draw(canvas)
    draw.text((52, 40), "系统输入与结果展示界面", fill="#17395f", font=pil_font(40, bold=True))
    draw.text((52, 92), "界面截图来自论文整理阶段的真实项目运行过程，保留了上传识别与在线手写两类入口。", fill="#6d829b", font=pil_font(20))
    ui_card(canvas, (48, 154, 876, 934), "图片上传识别与结果展示", "运行界面 A", UI_DIR / "upload_result_runtime.png", "#eef6ff")
    ui_card(canvas, (924, 154, 1752, 934), "在线手写板输入界面", "运行界面 B", UI_DIR / "sketch_board_runtime.png", "#fff4ea")
    result["ui_input"] = FIG_DIR / "figure_4_ui_input.png"
    canvas.save(result["ui_input"], quality=95)

    canvas = Image.new("RGB", (1800, 980), "#f8fafc")
    draw = ImageDraw.Draw(canvas)
    draw.text((52, 40), "勘误表推荐与历史记录界面", fill="#17395f", font=pil_font(40, bold=True))
    draw.text((52, 92), "勘误入口支持候选推荐与手动改写，历史页同步保留原识别结果、修正状态与导出内容。", fill="#6d829b", font=pil_font(20))
    ui_card(canvas, (48, 154, 1048, 934), "勘误表候选推荐与人工修正", "闭环界面 A", UI_DIR / "correction_modal_runtime.png", "#eefaf4")
    ui_card(canvas, (1096, 154, 1752, 934), "识别历史记录与纠错追踪", "闭环界面 B", UI_DIR / "history_runtime.png", "#fff7ee")
    result["ui_correction"] = FIG_DIR / "figure_4_ui_correction.png"
    canvas.save(result["ui_correction"], quality=95)
    return result


def make_metric_bar() -> Path:
    fp = plot_font(14)
    labels = ["V2 原始模型", "V3 改进模型"]
    seq = [70.56, 83.24]
    tok = [86.58, 92.69]
    x = [0, 1]
    plt.figure(figsize=(10.2, 5.8), facecolor="white")
    plt.bar([i - 0.17 for i in x], seq, width=0.34, color="#547da9", label="序列准确率")
    plt.bar([i + 0.17 for i in x], tok, width=0.34, color="#f08a24", label="符号准确率")
    for i, value in enumerate(seq):
        plt.text(i - 0.17, value + 0.6, f"{value:.2f}%", ha="center", fontproperties=fp)
    for i, value in enumerate(tok):
        plt.text(i + 0.17, value + 0.6, f"{value:.2f}%", ha="center", fontproperties=fp)
    plt.xticks(x, labels, fontproperties=fp)
    plt.yticks(fontproperties=fp)
    plt.ylim(60, 100)
    plt.ylabel("准确率（%）", fontproperties=fp)
    plt.title("V2 与 V3 模型性能对比", fontproperties=plot_font(18))
    plt.grid(axis="y", linestyle="--", alpha=0.22)
    plt.legend(prop=fp, frameon=False)
    out = FIG_DIR / "figure_5_metrics_bar.png"
    plt.tight_layout()
    plt.savefig(out, dpi=220, bbox_inches="tight", pad_inches=0.05)
    plt.close()
    return out


def make_seq_curve(metrics: dict) -> Path:
    plt.figure(figsize=(10.4, 5.8), facecolor="white")
    fp = plot_font(14)
    for label, color, limit in [("V2", "#557eac", 150), ("V3", "#e15a58", 200)]:
        rows = [(row["epoch"], row["val_seq_acc"] * 100) for row in metrics[label] if row["epoch"] <= limit]
        plt.plot([x for x, _ in rows], [y for _, y in rows], linewidth=2.8, color=color, label=f"{label} 验证集序列准确率")
    plt.xlabel("训练轮次（epoch）", fontproperties=fp)
    plt.ylabel("Sequence Accuracy（%）", fontproperties=fp)
    plt.title("验证集序列准确率变化曲线", fontproperties=plot_font(18))
    plt.xticks(fontproperties=fp)
    plt.yticks(fontproperties=fp)
    plt.grid(linestyle="--", alpha=0.22)
    plt.legend(prop=fp, frameon=False, loc="lower right")
    out = FIG_DIR / "figure_5_seq_curve.png"
    plt.tight_layout()
    plt.savefig(out, dpi=220, bbox_inches="tight", pad_inches=0.05)
    plt.close()
    return out


def make_runtime_donut(stats: dict) -> Path:
    values = [stats["history_source"].get("upload", 0), stats["history_source"].get("sketch", 0)]
    labels = ["图片上传识别", "在线手写识别"]
    colors = ["#6fa8dc", "#f6b26b"]
    fp = plot_font(13)
    fig, ax = plt.subplots(figsize=(6.8, 5.1), facecolor="white")
    _, texts, autotexts = ax.pie(
        values,
        labels=labels,
        colors=colors,
        autopct="%1.1f%%",
        startangle=90,
        wedgeprops={"width": 0.42, "edgecolor": "white", "linewidth": 2},
        pctdistance=0.8,
        textprops={"fontproperties": fp},
    )
    for text in autotexts:
        text.set_fontproperties(fp)
    ax.text(0, 0, f"{sum(values)}\n条记录", ha="center", va="center", fontproperties=plot_font(18))
    ax.set_title("识别记录来源分布", fontproperties=plot_font(18))
    out = FIG_DIR / "figure_5_runtime_donut.png"
    plt.tight_layout()
    plt.savefig(out, dpi=220, bbox_inches="tight", pad_inches=0.05)
    plt.close()
    return out


def make_correction_bar(stats: dict) -> Path:
    mapping = [
        ("builtin_formula", "系统内置勘误种子"),
        ("upload", "上传识别产生"),
        ("sketch", "手写识别产生"),
    ]
    labels = [item[1] for item in mapping]
    values = [stats["correction_source"].get(item[0], 0) for item in mapping]
    fp = plot_font(13)
    fig, ax = plt.subplots(figsize=(7.8, 5.0), facecolor="white")
    bars = ax.barh(labels, values, color=["#5b8bd9", "#79c4a4", "#f2aa62"])
    for bar, value in zip(bars, values):
        ax.text(bar.get_width() + 0.15, bar.get_y() + bar.get_height() / 2, str(value), va="center", fontproperties=fp)
    ax.set_xlabel("条目数量", fontproperties=fp)
    ax.set_title("勘误知识表来源构成", fontproperties=plot_font(18))
    ax.grid(axis="x", linestyle="--", alpha=0.22)
    for label in ax.get_xticklabels():
        label.set_fontproperties(fp)
    for label in ax.get_yticklabels():
        label.set_fontproperties(fp)
    out = FIG_DIR / "figure_5_correction_bar.png"
    plt.tight_layout()
    plt.savefig(out, dpi=220, bbox_inches="tight", pad_inches=0.05)
    plt.close()
    return out


def copy_support_assets() -> dict:
    assets = {}
    mapping = {
        "real_case_gallery": V2_FIG_DIR / "real_case_gallery.png",
        "real_case_stats": V2_FIG_DIR / "real_case_stats.png",
        "dataset_split": V2_FIG_DIR / "dataset_split_pie_cn.png",
    }
    for key, source in mapping.items():
        if source.exists():
            target = FIG_DIR / source.name
            shutil.copy2(source, target)
            assets[key] = target
    return assets


def generate_assets() -> dict:
    ensure_dirs()
    assets = copy_support_assets()
    assets.update(make_ui_collages())
    assets["architecture"] = make_architecture_figure()
    assets["metrics_bar"] = make_metric_bar()
    assets["seq_curve"] = make_seq_curve(load_training_metrics())
    stats = load_runtime_stats()
    assets["runtime_donut"] = make_runtime_donut(stats)
    assets["correction_bar"] = make_correction_bar(stats)
    return assets


def clear_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    for child in list(element):
        element.remove(child)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def insert_picture_after(paragraph: Paragraph, image_path: Path, width_cm: float) -> Paragraph:
    pic_para = insert_paragraph_after(paragraph)
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_para.add_run().add_picture(str(image_path), width=Cm(width_cm))
    return pic_para


def insert_table_after(doc: Document, paragraph: Paragraph, rows: list[list[str]]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    paragraph._p.addnext(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.cell(r_index, c_index)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def set_run_fonts(run, chinese: str = "宋体", western: str = "Times New Roman", size: float = 12.0, bold: bool | None = None) -> None:
    run.font.name = western
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), chinese)
    rfonts.set(qn("w:ascii"), western)
    rfonts.set(qn("w:hAnsi"), western)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def format_paragraph_runs(paragraph: Paragraph, chinese: str = "宋体", western: str = "Times New Roman", size: float = 12.0, bold: bool | None = None) -> None:
    if not paragraph.runs:
        paragraph.add_run("")
    for run in paragraph.runs:
        set_run_fonts(run, chinese=chinese, western=western, size=size, bold=bold)


def format_body(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    format_paragraph_runs(paragraph, chinese="宋体", western="Times New Roman", size=12)


def format_heading1(paragraph: Paragraph) -> None:
    paragraph.style = "Heading 1"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = not paragraph.text.strip().startswith("1 ")
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(18)
    paragraph.paragraph_format.line_spacing = 1.5
    format_paragraph_runs(paragraph, chinese="黑体", western="Times New Roman", size=18, bold=True)


def format_heading2(paragraph: Paragraph) -> None:
    paragraph.style = "Heading 2"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(9)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.line_spacing = 1.5
    format_paragraph_runs(paragraph, chinese="宋体", western="Times New Roman", size=14, bold=True)


def format_heading3(paragraph: Paragraph) -> None:
    paragraph.style = "Heading 3"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(9)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.line_spacing = 1.5
    format_paragraph_runs(paragraph, chinese="宋体", western="Times New Roman", size=12, bold=True)


def format_caption(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.5
    format_paragraph_runs(paragraph, chinese="宋体", western="Times New Roman", size=10.5)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("left", "top", "right", "bottom"):
        edge_data = kwargs.get(edge)
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if edge_data:
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ("val", "sz", "space", "color"):
                if key in edge_data:
                    element.set(qn("w:" + key), str(edge_data[key]))
        elif element is not None:
            tc_borders.remove(element)


def apply_three_line_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                left={"val": "nil"},
                right={"val": "nil"},
                top={"val": "nil"},
                bottom={"val": "nil"},
            )
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.first_line_indent = Cm(0)
                para.paragraph_format.line_spacing = 1.5
                format_paragraph_runs(para, chinese="宋体", western="Times New Roman", size=12)
    for cell in table.rows[0].cells:
        set_cell_border(
            cell,
            top={"val": "single", "sz": "12", "color": "000000"},
            bottom={"val": "single", "sz": "8", "color": "000000"},
            left={"val": "nil"},
            right={"val": "nil"},
        )
        for para in cell.paragraphs:
            format_paragraph_runs(para, chinese="宋体", western="Times New Roman", size=12, bold=True)
    for cell in table.rows[-1].cells:
        set_cell_border(
            cell,
            bottom={"val": "single", "sz": "12", "color": "000000"},
            left={"val": "nil"},
            right={"val": "nil"},
        )


def find_paragraph(doc: Document, startswith: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(startswith):
            return paragraph
    raise ValueError(f"Cannot find paragraph starting with: {startswith}")


def add_toc_field(paragraph: Paragraph) -> None:
    clear_paragraph(paragraph)
    r1 = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    r1._r.append(fld_begin)

    r2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    r2._r.append(instr)

    paragraph.add_run("目录将在 Word 中自动更新")

    r4 = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r4._r.append(fld_end)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.5
    format_paragraph_runs(paragraph, chinese="宋体", western="Times New Roman", size=12)


def extract_references() -> list[str]:
    source_doc = Document(str(SOURCE_DOC))
    refs = []
    started = False
    for paragraph in source_doc.paragraphs:
        text = paragraph.text.strip()
        if text == "参考文献":
            started = True
            continue
        if started and text:
            refs.append(text)
    return refs


def append_paragraph(doc: Document, text: str) -> Paragraph:
    para = doc.add_paragraph()
    para.add_run(text)
    return para


def apply_global_formatting(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if text == "目    录":
            format_heading1(paragraph)
            paragraph.paragraph_format.page_break_before = False
            continue
        if text == "参考文献":
            format_heading1(paragraph)
            continue
        if re.match(r"^\d+\.\d+\.\d+\s+\S+", text):
            format_heading3(paragraph)
        elif re.match(r"^\d+\.\d+\s+\S+", text):
            format_heading2(paragraph)
        elif re.match(r"^\d+\s{1,}\S+", text):
            format_heading1(paragraph)
        elif re.match(r"^[图表]\d+-\d+", text):
            format_caption(paragraph)
        elif re.match(r"^\[\d+\]", text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.hanging_indent = Cm(0.74)
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            format_paragraph_runs(paragraph, chinese="宋体", western="Times New Roman", size=12)
        else:
            format_body(paragraph)
    for index, table in enumerate(doc.tables):
        if index == 0:
            continue
        apply_three_line_table(table)


def build_document(assets: dict) -> None:
    doc = Document(str(BASE_DOC))

    p27 = doc.paragraphs[27]
    toc_title = doc.paragraphs[33]
    toc_field = doc.paragraphs[34]
    p_1_4 = find_paragraph(doc, "1.4 ")
    old_structure_para = find_paragraph(doc, "全文共分为六章。第一章介绍研究背景")
    p_4_1_2_body = find_paragraph(doc, "在线手写板使用 Canvas 构建书写区域")
    p_4_4 = find_paragraph(doc, "4.4 ")
    p_4_4_body_2 = find_paragraph(doc, "在此基础上，系统调用")
    p_4_4_body_3 = find_paragraph(doc, "进一步来看，")
    p_4_5 = find_paragraph(doc, "4.5 ")
    p_4_5_body_1 = find_paragraph(doc, "系统会将原始图像")
    p_4_4_fig = find_paragraph(doc, "图4-4 ")
    p_5_1 = find_paragraph(doc, "5.1 ")
    p_5_1_body_1 = find_paragraph(doc, "本文实验基于 CROHME 数据集开展")
    p_5_1_body_2 = find_paragraph(doc, "项目按 9:1 的比例划分训练集")
    p_5_1_fig = find_paragraph(doc, "图5-1 ")

    if "方    式" in p27.text:
        set_paragraph_text(p27, p27.text.replace("方    式", "方式"))
    set_paragraph_text(old_structure_para, "")

    para_4_3_4 = find_paragraph(doc, "系统在工程上兼容 LSTM 解码器")
    para_4_3_5 = find_paragraph(doc, "V3 模型的另一项关键改进是 Coverage Attention")
    set_paragraph_text(para_4_3_4, para_4_3_4.text.replace("[19] [36]", "[19][36]"))
    set_paragraph_text(para_4_3_5, para_4_3_5.text.replace("连锁影响,这", "连锁影响，这"))

    set_paragraph_text(toc_title, "目    录")
    add_toc_field(toc_field)

    set_paragraph_text(p_1_4, "1.4 系统总体架构与技术路线")
    p_1_4_body = insert_paragraph_after(
        p_1_4,
        "结合项目实现过程可以看到，本文的研究思路并不是把“模型训练”和“系统开发”割裂开来，而是围绕课堂公式录入这一具体场景，把输入采集、图像预处理、序列识别、多格式输出、历史留存与人工纠错组织为同一条工作链路。与只关注离线精度的纯算法实验相比，这种架构更强调前后端协同和结果可复用性，也更贴合本科毕业设计对可展示性与完整性的要求。",
    )
    fig_1 = insert_picture_after(p_1_4_body, assets["architecture"], 15.8)
    cap_1 = insert_paragraph_after(fig_1, "图1-1 系统总体架构与技术路线图")
    p_1_5 = insert_paragraph_after(cap_1, "1.5 论文结构安排")
    insert_paragraph_after(
        p_1_5,
        "全文共分为六章。第一章介绍课题背景、研究现状以及本文的总体技术路线；第二章归纳手写数学公式识别涉及的相关理论与关键技术；第三章从需求分析、系统角色与模块划分等角度给出总体设计；第四章结合项目代码，对前端输入、预处理、识别模型、结果转换、历史记录与勘误闭环等关键实现进行详细说明；第五章基于训练日志、实际运行界面和项目数据文件，对模型性能、系统功能和真实样例进行综合测试与分析；第六章总结全文工作，并讨论后续仍可继续完善的方向。",
    )

    p_4_1_3 = insert_paragraph_after(p_4_1_2_body, "4.1.3 统一请求与结果回传链路")
    insert_paragraph_after(
        p_4_1_3,
        "从后端组织方式看，图片上传与在线手写虽然入口不同，但在进入推理前都会被转成统一的图像对象，并共用同一套预处理与结果封装逻辑。app.py 中的 _build_result_payload() 会一次性返回 LaTeX 文本、Word Linear、MathML、历史记录编号、公式渲染图地址和勘误推荐项。这样做的好处是前端页面不必为不同输入方式分别维护两套展示结构，用户也能在识别完成后立即复制、导出或进入纠错流程，界面行为更连贯。",
    )

    insert_paragraph_after(p_4_4, "4.4.1 多格式转换与公式渲染")
    p_4_4_2 = insert_paragraph_after(p_4_4_body_2, "4.4.2 表达式化简与方程求解")
    insert_paragraph_after(
        p_4_4_2,
        "项目并未把识别结果停留在“字符串输出”层面，而是继续借助 equation_calculator.py 和 solve_equation_file.py 完成表达式解析、化简与方程求解。对于教学场景而言，这一步的意义在于把识别系统从单纯录入工具扩展为可继续计算的辅助工具。换句话说，识别后的公式不仅能看、能复制，还能进一步参与推导与验证，这会显著提升系统在作业批改、板书整理和课堂演示中的实用价值。",
    )
    set_paragraph_text(
        p_4_4_body_3,
        "进一步来看，格式转换并不是孤立步骤，而是后续勘误推荐的重要前置条件。只有先把识别结果尽量归一到稳定的内部表达形式，系统才能较可靠地利用错误公式本身、扩展匹配词和相似度评分筛选候选项。对于常见教学公式，这种“标准化表达 + 轻量级候选推荐”的组合方式比单纯字符串比对更稳妥，也更符合本科项目在实现成本与实用性之间求平衡的定位。",
    )

    set_paragraph_text(p_4_5, "4.5 历史记录与勘误闭环实现")
    insert_paragraph_after(p_4_5, "4.5.1 历史记录数据组织")
    p_4_5_2 = insert_paragraph_after(p_4_5_body_1, "4.5.2 勘误知识表与人工回写")
    insert_paragraph_after(
        p_4_5_2,
        "从数据库实现细节看，record_correction() 会把新的错误公式与正确公式配对写入 correction_table；如果同类修正已经存在，则只累计使用次数并刷新更新时间。与此同时，update_history_correction() 会把原识别结果转存到 original_recognized_text，再用修正后的公式覆盖 recognized_text。这样一来，系统既保留了“机器第一次给出的答案”，也保留了“用户最终确认的答案”，为后续分析哪些公式经常出错、哪些候选项更常被采用提供了可追踪依据。",
    )

    fig_4_5 = insert_picture_after(p_4_4_fig, assets["ui_input"], 16.0)
    cap_4_5 = insert_paragraph_after(fig_4_5, "图4-5 系统输入与结果展示界面")
    fig_4_6 = insert_picture_after(cap_4_5, assets["ui_correction"], 16.0)
    insert_paragraph_after(fig_4_6, "图4-6 勘误表推荐与历史记录界面")

    set_paragraph_text(p_5_1, "5.1 实验数据与评价设置")
    insert_paragraph_after(p_5_1, "5.1.1 数据来源与日志说明")
    set_paragraph_text(p_5_1_body_1, "本文实验基于 CROHME 数据集开展，数据来源覆盖 2014、2016、2019 和 train 等划分，图像高度统一为 64，最大宽度限制为 512，最大序列长度设置为 256。CROHME 作为手写数学公式识别领域使用最广的公开基准之一，为不同方法之间的横向比较提供了较稳定的评价平台[8][9][10][11]。")
    set_paragraph_text(p_5_1_body_2, "除公开数据集外，本文在论文整理阶段还结合项目中的训练日志、识别历史记录、勘误知识表以及真实运行界面进行补充分析。这样处理的原因在于，本项目的目标并不仅是给出一组离线精度数字，更希望说明一个课堂导向的 Web 系统在真实使用链路中是否可用、是否容易修正错误、以及是否具备持续积累经验的能力。")
    p_5_1_2 = insert_paragraph_after(p_5_1_body_2, "5.1.2 评价指标与工程验证口径")
    insert_paragraph_after(
        p_5_1_2,
        "模型性能仍以 sequence accuracy 和 token accuracy 为主。前者要求整条公式完全正确，更接近真实使用时“能否直接拿去用”的判断标准；后者反映局部符号预测质量，适合观察模型在细粒度层面的提升情况。除此之外，本文还增加了两类工程化验证口径：一类是系统功能是否能够完整跑通，包括登录、上传、手写、导出、历史回看与勘误回写；另一类是项目数据文件中已经积累的运行记录，它能够从侧面反映系统在不同输入方式和纠错场景下的实际使用特征[6]。",
    )
    set_paragraph_text(p_5_1_fig, "图5-1 数据集划分与训练来源示意图")

    append_paragraph(doc, "5.2 模型性能对比分析")
    append_paragraph(doc, "从项目保留的训练日志与检查点指标来看，V2 和 V3 模型的主体框架保持一致，差异主要集中在编码增强、对齐建模、数据增强和解码策略等几个关键环节。正因为改动具有明确边界，因此两版模型的对比更能说明“具体工程改进是否带来了可观收益”，而不是被其他无关变量干扰。")
    p = append_paragraph(doc, "表5-1 V2 与 V3 模型关键配置对比")
    insert_table_after(
        doc,
        p,
        [
            ["比较项", "V2 原始模型", "V3 改进模型"],
            ["编码端结构", "基础 CNN 编码器", "更深 CNN 编码器，并引入 SE 模块"],
            ["注意力建模", "常规注意力机制", "加入 Coverage Attention"],
            ["训练增强", "常规图像预处理", "弹性形变、CutOut、随机擦除等组合增强"],
            ["解码方式", "单路径贪心解码", "Beam Search 多候选解码"],
            ["整体目标", "实现基础识别闭环", "提升结构复杂公式的整体稳定性"],
        ],
    )
    append_paragraph(doc, "结合训练日志整理得到的最佳指标显示，V2 模型的 sequence accuracy 为 70.56%，token accuracy 为 86.58%；V3 模型分别提升到 83.24% 和 92.69%。如果只看数值，二者差距已经比较明显；若结合项目场景理解，这种提升更重要的意义在于它减少了“公式大体看起来像对了，但关键结构仍然错位”的情况。")
    p = append_paragraph(doc, "表5-2 不同模型识别性能对比")
    insert_table_after(
        doc,
        p,
        [
            ["指标", "V2", "V3", "提升幅度"],
            ["Sequence Accuracy", "70.56%", "83.24%", "+12.68%"],
            ["Token Accuracy", "86.58%", "92.69%", "+6.11%"],
        ],
    )
    fig = append_paragraph(doc, "")
    fig.add_run().add_picture(str(assets["metrics_bar"]), width=Cm(14.8))
    fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    append_paragraph(doc, "图5-2 V2 与 V3 模型性能对比图")
    fig = append_paragraph(doc, "")
    fig.add_run().add_picture(str(assets["seq_curve"]), width=Cm(15.2))
    fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    append_paragraph(doc, "图5-3 验证集序列准确率变化曲线")
    append_paragraph(doc, "从曲线走势看，V3 不只是最终峰值更高，在中后期训练阶段也表现出更稳定的平台区间。这说明 SE 模块、Coverage 建模与更强的数据增强并非只在个别 epoch 上“碰巧有效”，而是整体改善了模型对复杂结构和分布扰动的适应能力[20][37][49]。")

    append_paragraph(doc, "5.3 系统运行与功能测试")
    stats = load_runtime_stats()
    append_paragraph(doc, f"在系统运行层面，项目数据文件中已经累计保存 {stats['history_total']} 条识别记录，其中图片上传识别 {stats['history_source'].get('upload', 0)} 条，在线手写识别 {stats['history_source'].get('sketch', 0)} 条。勘误知识表当前共 {stats['correction_total']} 条，其中 {stats['correction_source'].get('builtin_formula', 0)} 条为系统初始化的常见公式种子，另有 {stats['correction_source'].get('upload', 0) + stats['correction_source'].get('sketch', 0)} 条来自真实使用过程中的人工修正。虽然这些数据规模还不算大，但它们足以说明勘误闭环已经不再停留在接口设计层面，而是开始在项目内部产生可复用的经验沉淀。")
    append_paragraph(doc, "结合本次实际运行测试，可以确认登录认证、图片上传、在线手写、LaTeX/Word/PNG 导出、历史记录查看和勘误弹窗等核心功能均能正常工作。尤其是在上传二次公式样例后，系统可以同时返回公式渲染图、LaTeX 文本、Word Linear 结果以及多个勘误推荐项，这一过程与第四章对实现链路的描述能够互相印证。")
    p = append_paragraph(doc, "表5-3 系统主要功能测试结果")
    insert_table_after(
        doc,
        p,
        [
            ["测试项", "验证方式", "结果"],
            ["登录与角色校验", "管理员身份登录并进入首页", "通过"],
            ["图片上传识别", "上传二次公式样例并获取识别结果", "通过"],
            ["在线手写输入", "打开画板、写入笔迹并触发识别流程", "通过"],
            ["多格式导出", "复制 LaTeX、Word 和 PNG 结果", "通过"],
            ["历史记录管理", "查看、分页与复核已识别条目", "通过"],
            ["勘误表闭环", "打开候选修正弹窗并提交正确公式", "通过"],
        ],
    )
    fig = append_paragraph(doc, "")
    fig.add_run().add_picture(str(assets["runtime_donut"]), width=Cm(11.6))
    fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    append_paragraph(doc, "图5-4 识别记录来源分布图")
    fig = append_paragraph(doc, "")
    fig.add_run().add_picture(str(assets["correction_bar"]), width=Cm(12.4))
    fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    append_paragraph(doc, "图5-5 勘误知识表来源构成图")

    append_paragraph(doc, "5.4 真实样例与界面运行分析")
    append_paragraph(doc, "为了让测试结果更接近答辩现场的展示方式，本文继续保留 6 组真实样例的现场复测结果。这部分样例覆盖一次函数、根式、二次公式、三角形面积公式、三角函数分式和复杂分式等类型。测试流程与系统真实调用方式一致，即先进行预处理，再送入当前模型推理，最后结合人工判断给出“正确、部分正确、错误”三类评价。")
    if "real_case_gallery" in assets:
        fig = append_paragraph(doc, "")
        fig.add_run().add_picture(str(assets["real_case_gallery"]), width=Cm(15.4))
        fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
        append_paragraph(doc, "图5-6 真实样例现场测试结果图")
    real_cases = load_real_cases()
    p = append_paragraph(doc, "表5-4 真实样例现场测试结果汇总")
    rows = [["样例", "公式类型", "人工评价"]]
    for item in real_cases:
        rows.append([item["name"], item["kind"], item["judge"]])
    insert_table_after(doc, p, rows)
    if "real_case_stats" in assets:
        fig = append_paragraph(doc, "")
        fig.add_run().add_picture(str(assets["real_case_stats"]), width=Cm(11.6))
        fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
        append_paragraph(doc, "图5-7 真实样例人工评估统计图")
    append_paragraph(doc, "从复测结果看，系统对一次函数、简单根式、标准二次公式以及常见三角面积公式已经表现出较好的可用性；真正容易失稳的，仍然是层次更深、分式嵌套更复杂、或者局部连笔比较明显的样例。换句话说，当前模型在“常见教学公式”层面已经具备展示价值，但在更复杂结构下仍需要依赖后续纠错或进一步训练来保证结果可靠。")

    append_paragraph(doc, "5.5 结果讨论与不足")
    append_paragraph(doc, "综合模型指标与系统运行情况可以发现，V3 相比 V2 的优势并不仅体现在离线精度提升上，更体现在应用过程中的容错空间更大。对于用户来说，识别结果越接近正确答案，后续需要人工调整的幅度就越小；而当识别仍然出错时，勘误表推荐又能帮助用户更快完成修正。正是“模型改进 + 人机协同反馈”这两个层面叠加，才让系统整体可用性有了比较明显的提升。")
    append_paragraph(doc, "当然，本项目仍然存在几方面不足。第一，当前工程运行记录主要来自项目开发与论文整理阶段，真实课堂规模下的长周期使用数据还不充分；第二，模型面对复杂分式、非常规连写和拍照噪声较重的样本时，依旧可能出现结构重复或错位问题[52][54]；第三，勘误知识表目前更适合作为轻量级经验库使用，若后续希望面向更多用户共享，还需要继续完善审核机制、权限控制与质量筛选策略[48]。")

    append_paragraph(doc, "6  总结与展望")
    append_paragraph(doc, "6.1 全文总结")
    append_paragraph(doc, "本文围绕课堂场景下的手写数学公式录入需求，设计并实现了一套基于 Flask 的 Web 端识别系统。系统支持图片上传与在线手写两类输入方式，并在后端完成统一预处理、端到端识别、多格式转换、公式计算、历史记录管理和勘误表闭环纠错等功能。从项目完成度看，当前版本已经不再只是一个能跑通模型推理的原型，而是一套具备输入、识别、展示、修正与积累能力的完整应用链路。")
    append_paragraph(doc, "论文整理过程中进一步结合代码、运行界面与项目数据文件，对第四章和第五章进行了更细致的补充。新的章节内容不仅说明了关键模块“是什么”，也尽量回答了它们“为什么这样设计、真实运行时表现如何、当前边界在哪里”。这种写法更能体现本科毕业设计面向实际问题解决的特点，也更方便后续答辩时从界面、流程和数据三个层面进行说明。")
    append_paragraph(doc, "6.2 展望")
    append_paragraph(doc, "后续工作仍有较大的完善空间。首先，可以继续扩充真实纸面样本与课堂板书样本，缩小训练数据与实际应用场景之间的分布差距；其次，可以围绕复杂分式、长公式和多层上下标继续加强结构约束建模，提升极端样例下的稳定性；再次，可以把勘误表从静态候选推荐逐步发展为更细致的主动学习或人工审核机制，让纠错信息更有效地反哺系统；最后，若未来希望服务更大范围的教学使用，还需要从权限管理、数据共享与日志分析等角度继续完善整个系统。")

    append_paragraph(doc, "参考文献")
    for ref in extract_references():
        append_paragraph(doc, ref)

    apply_global_formatting(doc)
    doc.save(str(OUTPUT_DOC))


def build_note() -> None:
    content = """# 论文4.0 修改说明

## 本轮重点修改

1. 将原 `1.4 论文结构安排` 调整为 `1.4 系统总体架构与技术路线`，新增 `图1-1`，并补写新的 `1.5 论文结构安排`。
2. 在第四章新增 `4.1.3 统一请求与结果回传链路`，补强前后端如何共用同一套结果封装逻辑。
3. 在第四章补入 `4.4.1 多格式转换与公式渲染`、`4.4.2 表达式化简与方程求解`，增强结果后处理部分的工程说明。
4. 将 `4.5 历史记录管理实现` 调整为 `4.5 历史记录与勘误闭环实现`，新增 `4.5.1` 与 `4.5.2`，把勘误表功能写成完整闭环。
5. 新增 `图4-5`、`图4-6`，使用真实项目运行截图展示输入界面、勘误弹窗和历史记录。
6. 重建第五章内容，增加 `5.1.1`、`5.1.2`，补写模型性能、系统运行、真实样例和结果讨论等部分。
7. 新增 `表5-1` 至 `表5-4`，并将正文业务表统一处理为三线表。
8. 新增 `图5-2` 至 `图5-7`，包括性能柱状图、训练曲线、运行记录分布图、勘误来源构成图和真实样例图。
9. 新增第六章总结与展望，并将 `3.4` 版本中的参考文献完整移入当前稿件。
10. 统一正文标题、二三级标题、正文段落、图表标题和参考文献版式，便于后续在 Word 中刷新目录。
"""
    NOTE_FILE.write_text(content, encoding="utf-8")


def main() -> None:
    assets = generate_assets()
    build_document(assets)
    build_note()
    print(f"Generated: {OUTPUT_DOC}")
    print(f"Assets: {FIG_DIR}")
    print(f"Note: {NOTE_FILE}")


if __name__ == "__main__":
    main()
